"""Leitura OOXML (openpyxl) de backups Sólides — colaboradores, solicitações,
crosswalk, notas, comentários, habilidades e PDIs.

Único módulo autorizado a importar ``openpyxl`` (research R1; 010/011/013/014).
Colunas: ``contracts/column-mapping-contract.md``; pré-condições:
``contracts/import-command-contract.md``.

015: ``parse_colaboradores_admission_xlsx`` lê ``Data admissão`` sem mutar
``ColaboradorRow`` / caminho do importer 010 (datas via ``parse_legacy_date``).

Falhas fatais pré-persistência (arquivo ausente/ilegível, OOXML inválido,
colunas obrigatórias ausentes) → ``LegacyParseError`` (exit 1; zero writes).
"""

from __future__ import annotations

import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.competencies.services.catalog_import.normalize import display_name

# --- Colunas (contrato) -------------------------------------------------------

COLABORADORES_REQUIRED_COLUMNS: tuple[str, ...] = ("Nome",)

COLABORADORES_COLUMNS: tuple[str, ...] = (
    "Nome",
    "E-mail empresarial",
    "E-mail",
    "E-mail pessoal",
    "Data demissão",
    "Cargo",
    "Cargo ID",
    "Departamento",
    "Superior direto id",
)

# Backfill 015 — parser dedicado (MUST NOT mutar ColaboradorRow / importer 010).
ADMISSION_BACKFILL_REQUIRED_COLUMNS: tuple[str, ...] = ("Data admissão",)

ADMISSION_BACKFILL_COLUMNS: tuple[str, ...] = (
    "Nome",
    "E-mail empresarial",
    "E-mail",
    "E-mail pessoal",
    "Identificador",
    "Data admissão",
)

AVALIACOES_REQUIRED_COLUMNS: tuple[str, ...] = (
    "Nome Avaliado",
    "Identificador Avaliado",
)

AVALIACOES_HEADERS_REQUIRED_COLUMNS: tuple[str, ...] = (
    "Identificador",
    "Identificador Solicitação",
    "Identificador Avaliado",
    "Nome Avaliado",
    "Nome Avaliador",
)

SOLICITACOES_REQUIRED_COLUMNS: tuple[str, ...] = (
    "Identificador",
    "Nome",
    "Iniciada em",
    "Terminada em",
)

SOLICITACOES_OPTIONAL_COLUMNS: tuple[str, ...] = ("Status",)

NOTAS_REQUIRED_COLUMNS: tuple[str, ...] = (
    "Identificador",
    "Identificador Avaliação",
    "Nome Avaliador",
    "Nome Avaliado",
    "Identificador Habilidade",
    "habilidade",
    "Fator no Momento",
    "Nota",
)

COMENTARIOS_REQUIRED_COLUMNS: tuple[str, ...] = (
    "Identificador",
    "Identificador Avaliador",
    "Nome Avaliador",
    "Nome Avaliado",
    "Comentário",
    "Criado em",
)

HABILIDADES_REQUIRED_COLUMNS: tuple[str, ...] = (
    "Identificador",
    "Habilidade",
    "Grupo",
)

PDI_REQUIRED_COLUMNS: tuple[str, ...] = (
    "Nome",
    "Título do PDI",
    "Status",
    "Objetivo",
    "Situação Atual",
    "Situação Desejada",
    "Data de Entrega",
)

PDI_OPTIONAL_COLUMNS: tuple[str, ...] = (
    "Criado em",
    "Identificador Solicitação",
    "Identificador",
    "Identificador Avaliado",
)


class LegacyParseError(Exception):
    """Erro fatal de pré-condição de parse (arquivo, OOXML, colunas).

    Propagado para o management command → exit ``1``, sem escrita no DB.
    """


@dataclass(frozen=True)
class ColaboradorRow:
    """Linha normalizada de ``backup_colaboradores`` (sem PII extra).

    ``data_demissao`` permanece cru (datetime / serial / ISO / None) para
    ``dates.py`` (T006). IDs já como string canônica (sem ``.0`` de float).
    """

    linha: int
    nome: str
    email_empresarial: str
    email: str
    email_pessoal: str
    data_demissao: Any
    cargo: str
    cargo_id: str
    departamento: str
    superior_direto_id: str


@dataclass(frozen=True)
class ColaboradorAdmissionRow:
    """Linha mínima para backfill de ``data_entrada`` (015 — US4).

    ``data_admissao`` permanece **cru** (datetime / serial Excel / ISO /
    None) para ``dates.parse_legacy_date`` via
    ``admission_backfill.resolve.parse_data_admissao``. ``identificador``
    já canônico. Demais colunas do backup são ignoradas — **não** reabre
    o caminho de persistência do importer 010.
    """

    linha: int
    nome: str
    email_empresarial: str
    email: str
    email_pessoal: str
    identificador: str
    data_admissao: Any


@dataclass(frozen=True)
class SolicitacaoRow:
    """Linha normalizada de ``backup_solicitacoes_*`` (sem persistência).

    ``nome``, ``iniciada_em`` e ``terminada_em`` permanecem crus (texto /
    serial Excel / datetime / None) para ``dates.normalize_ciclo_nome`` e
    ``parse_legacy_date``. ``identificador`` já canônico via
    ``canonicalize_id``. ``status`` é opcional (contrato: sempre → encerrado).
    """

    linha: int
    identificador: str
    nome: Any
    iniciada_em: Any
    terminada_em: Any
    status: str


@dataclass(frozen=True)
class AvaliacaoCrosswalkRow:
    """Par Nome Avaliado → Identificador Avaliado para crosswalk (R8)."""

    linha: int
    nome_avaliado: str
    identificador_avaliado: str


@dataclass(frozen=True)
class AvaliacaoHeaderRow:
    """Linha de cabeçalho de ``backup_avaliacoes_*`` (pré-agregação).

    IDs já canônicos. Nomes via ``display_name`` para autoavaliação /
    fallback (``canonical_key``). Demais colunas do backup são ignoradas.
    """

    linha: int
    identificador: str
    identificador_solicitacao: str
    identificador_avaliado: str
    nome_avaliado: str
    nome_avaliador: str


@dataclass(frozen=True)
class NotaRow:
    """Linha de ``backup_notas_avaliacoes_*`` (sem persistência).

    IDs já canônicos. Nomes via ``display_name``. ``fator_no_momento`` e
    ``nota`` permanecem crus (número / texto / None) para validação na
    fase de snapshots. Dump 2026-06-24 **não** tem coluna de nível —
    não parsear / não inventar.
    """

    linha: int
    identificador: str
    identificador_avaliacao: str
    nome_avaliador: str
    nome_avaliado: str
    identificador_habilidade: str
    habilidade: str
    fator_no_momento: Any
    nota: Any


@dataclass(frozen=True)
class ComentarioRow:
    """Linha de ``backup_comentarios_avaliacoes_*`` (sem persistência).

    ``identificador`` é o ID de **avaliação** (não do comentário).
    ``criado_em`` permanece cru para ``dates.parse_legacy_datetime``.
    ``comentario`` via ``display_name``; não vai para o relatório.
    """

    linha: int
    identificador: str
    identificador_avaliador: str
    nome_avaliador: str
    nome_avaliado: str
    comentario: str
    criado_em: Any


@dataclass(frozen=True)
class HabilidadeRow:
    """Linha de ``backup_habilidades_*`` (opcional; só se ``--habilidades``).

    Completa ``nome``/``tipo`` de extras criadas como FK de nota (R11).
    **Não** é a matriz ``backup_habilidades_cargo_*``.
    """

    linha: int
    identificador: str
    habilidade: str
    grupo: str


@dataclass(frozen=True)
class PdiRow:
    """Linha de ``backup_pdi_*`` (sem persistência; 014).

    Textos e datas permanecem **crus** (datetime / serial Excel / ISO /
    None) para ``dates.py`` e o domínio em ``pdi.services.legacy_import``.
    IDs opcionais já canônicos via ``canonicalize_id``.
    """

    linha: int
    nome: Any
    titulo: Any
    status: Any
    objetivo: Any
    situacao_atual: Any
    situacao_desejada: Any
    data_entrega: Any
    criado_em: Any
    identificador_solicitacao: str
    identificador_pessoa: str


@dataclass(frozen=True)
class ParsedColaboradores:
    """Resultado do parse de colaboradores (sem persistência)."""

    rows: tuple[ColaboradorRow, ...]
    path: str


@dataclass(frozen=True)
class ParsedColaboradoresAdmission:
    """Resultado do parse de colaboradores só para backfill 015."""

    rows: tuple[ColaboradorAdmissionRow, ...]
    path: str


@dataclass(frozen=True)
class ParsedSolicitacoes:
    """Resultado do parse de solicitações (sem persistência)."""

    rows: tuple[SolicitacaoRow, ...]
    path: str


@dataclass(frozen=True)
class ParsedAvaliacoesCrosswalk:
    """Resultado do parse mínimo de avaliações (crosswalk only)."""

    rows: tuple[AvaliacaoCrosswalkRow, ...]
    path: str


@dataclass(frozen=True)
class ParsedAvaliacoesHeaders:
    """Resultado do parse de cabeçalhos de avaliação (sem agregação)."""

    rows: tuple[AvaliacaoHeaderRow, ...]
    path: str


@dataclass(frozen=True)
class ParsedNotas:
    """Resultado do parse de notas por competência (sem persistência)."""

    rows: tuple[NotaRow, ...]
    path: str


@dataclass(frozen=True)
class ParsedComentarios:
    """Resultado do parse de comentários qualitativos (sem persistência)."""

    rows: tuple[ComentarioRow, ...]
    path: str


@dataclass(frozen=True)
class ParsedHabilidades:
    """Resultado do parse opcional de ``backup_habilidades_*``."""

    rows: tuple[HabilidadeRow, ...]
    path: str


@dataclass(frozen=True)
class ParsedPdi:
    """Resultado do parse de ``backup_pdi_*`` (sem persistência)."""

    rows: tuple[PdiRow, ...]
    path: str


def _assert_readable_file(path: Path) -> None:
    """Garante que ``path`` existe e é um arquivo legível.

    Raises:
        LegacyParseError: ausente, não é arquivo, ou ilegível (T025).
    """
    if not path.exists():
        raise LegacyParseError(f"Arquivo não encontrado: {path}")
    if not path.is_file():
        raise LegacyParseError(f"Path não é um arquivo legível: {path}")
    try:
        with path.open("rb") as fh:
            fh.read(1)
    except OSError as exc:
        raise LegacyParseError(f"Arquivo ilegível: {path} ({exc})") from exc


def validate_source_paths(
    colaboradores_path: str | Path,
    avaliacoes_path: str | Path | None = None,
) -> tuple[Path, Path | None]:
    """Pré-condição: path colaboradores (e avaliações, se informado) existem.

    Valida os paths **antes** de ler conteúdo OOXML, para falhar cedo
    sem escrita. Se ambos falharem, a mensagem agrega os dois.

    Returns:
        ``(colaboradores_path, avaliacoes_path)`` resolvidos como ``Path``.
        ``avaliacoes_path`` é ``None`` quando não informado / vazio.

    Raises:
        LegacyParseError: um ou ambos os paths inválidos (T025 — fatal).
    """
    colaboradores = Path(colaboradores_path)
    avaliacoes: Path | None = None
    if avaliacoes_path is not None and str(avaliacoes_path).strip():
        avaliacoes = Path(avaliacoes_path)

    errors: list[str] = []
    for path in (colaboradores, avaliacoes):
        if path is None:
            continue
        try:
            _assert_readable_file(path)
        except LegacyParseError as exc:
            errors.append(str(exc))
    if errors:
        raise LegacyParseError("; ".join(errors))
    return colaboradores, avaliacoes


def _header_index(header_row: tuple[Any, ...]) -> dict[str, int]:
    """Mapeia nome de coluna (strip) → índice; primeira ocorrência vence."""
    index: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        name = str(cell).strip()
        if name and name not in index:
            index[name] = idx
    return index


def _require_columns(
    path: Path,
    header_index: dict[str, int],
    required: tuple[str, ...],
) -> None:
    missing = [col for col in required if col not in header_index]
    if missing:
        raise LegacyParseError(
            f"Colunas obrigatórias ausentes em {path}: {', '.join(missing)}"
        )


def _cell(row: tuple[Any, ...], header_index: dict[str, int], column: str) -> Any:
    idx = header_index.get(column)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _as_text(value: Any) -> str:
    """Texto display: None→''; floats inteiros sem '.0'; demais via display_name."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return display_name(str(value))
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return display_name(str(value))
    return display_name(str(value))


def canonicalize_id(value: Any) -> str:
    """Identificador Sólides canônico (contrato §Identificadores).

    ``123.0`` (float/int integral) → ``'123'``; demais → ``strip(str(value))``.
    Evita duplicatas ``'123'`` vs ``'123.0'`` em ``solides_id``.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).strip()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip()
    return str(value).strip()


def _as_id(value: Any) -> str:
    """Alias interno — preferir ``canonicalize_id`` na API pública."""
    return canonicalize_id(value)


def _load_workbook(path: Path):
    """Abre OOXML; falha de formato → ``LegacyParseError`` (T025)."""
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except InvalidFileException as exc:
        raise LegacyParseError(
            f"Arquivo OOXML ilegível (não é Excel 2007+ válido): {path}"
        ) from exc
    except (zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise LegacyParseError(
            f"Arquivo OOXML ilegível (não é Excel 2007+ válido): {path}"
        ) from exc
    except OSError as exc:
        raise LegacyParseError(f"Arquivo ilegível: {path} ({exc})") from exc


def _load_sheet_rows(path: Path) -> list[tuple[Any, ...]]:
    """Carrega a primeira planilha como lista de tuplas (header + dados).

    Raises:
        LegacyParseError: arquivo ausente/ilegível, OOXML inválido ou
            planilha ausente (T025 — fatal, zero writes).
    """
    _assert_readable_file(path)
    workbook = _load_workbook(path)
    try:
        try:
            sheet = workbook.active
            if sheet is None:
                raise LegacyParseError(f"Planilha ausente em {path}")
            rows: list[tuple[Any, ...]] = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(tuple(row))
            return rows
        except LegacyParseError:
            raise
        except (
            OSError,
            ValueError,
            KeyError,
            zipfile.BadZipFile,
            zipfile.LargeZipFile,
        ) as exc:
            raise LegacyParseError(
                f"Arquivo OOXML ilegível: {path} ({exc})"
            ) from exc
    finally:
        workbook.close()


def parse_colaboradores_xlsx(path: str | Path) -> ParsedColaboradores:
    """Lê ``backup_colaboradores_*.xlsx`` e valida coluna obrigatória ``Nome``.

    Colunas PII (CPF, RG, banco, etc.) são ignoradas — não entram na estrutura.
    Linhas sem ``Nome`` são descartadas. ``Data demissão`` permanece crua.

    Raises:
        LegacyParseError: arquivo / OOXML / colunas (T025 — fatal).
    """
    path = Path(path)
    raw_rows = _load_sheet_rows(path)
    if not raw_rows:
        raise LegacyParseError(f"Planilha vazia (sem header): {path}")

    header_index = _header_index(raw_rows[0])
    _require_columns(path, header_index, COLABORADORES_REQUIRED_COLUMNS)

    result: list[ColaboradorRow] = []
    for excel_row, row in enumerate(raw_rows[1:], start=2):
        nome = _as_text(_cell(row, header_index, "Nome"))
        if not nome:
            continue
        result.append(
            ColaboradorRow(
                linha=excel_row,
                nome=nome,
                email_empresarial=_as_text(
                    _cell(row, header_index, "E-mail empresarial")
                ),
                email=_as_text(_cell(row, header_index, "E-mail")),
                email_pessoal=_as_text(_cell(row, header_index, "E-mail pessoal")),
                data_demissao=_cell(row, header_index, "Data demissão"),
                cargo=_as_text(_cell(row, header_index, "Cargo")),
                cargo_id=_as_id(_cell(row, header_index, "Cargo ID")),
                departamento=_as_text(_cell(row, header_index, "Departamento")),
                superior_direto_id=_as_id(
                    _cell(row, header_index, "Superior direto id")
                ),
            )
        )

    return ParsedColaboradores(rows=tuple(result), path=str(path))


def parse_colaboradores_admission_xlsx(
    path: str | Path,
) -> ParsedColaboradoresAdmission:
    """Lê ``backup_colaboradores_*`` só para backfill de ``data_entrada`` (015).

    Coluna obrigatória: ``Data admissão``. Lê e-mails, ``Identificador`` e
    ``Nome`` quando presentes. ``Data admissão`` permanece crua para
    ``parse_legacy_date`` (não grava; não toca importer 010).

    Linhas sem chave de match (nenhum e-mail e sem ``Identificador``) são
    descartadas. Células vazias de data **não** são erro de parse.

    Raises:
        LegacyParseError: arquivo / OOXML / coluna ``Data admissão`` ausente.
    """
    path = Path(path)
    raw_rows = _load_sheet_rows(path)
    if not raw_rows:
        raise LegacyParseError(f"Planilha vazia (sem header): {path}")

    header_index = _header_index(raw_rows[0])
    _require_columns(path, header_index, ADMISSION_BACKFILL_REQUIRED_COLUMNS)

    result: list[ColaboradorAdmissionRow] = []
    for excel_row, row in enumerate(raw_rows[1:], start=2):
        email_empresarial = _as_text(
            _cell(row, header_index, "E-mail empresarial")
        )
        email = _as_text(_cell(row, header_index, "E-mail"))
        email_pessoal = _as_text(_cell(row, header_index, "E-mail pessoal"))
        identificador = canonicalize_id(_cell(row, header_index, "Identificador"))
        if not email_empresarial and not email and not email_pessoal and not identificador:
            continue
        result.append(
            ColaboradorAdmissionRow(
                linha=excel_row,
                nome=_as_text(_cell(row, header_index, "Nome")),
                email_empresarial=email_empresarial,
                email=email,
                email_pessoal=email_pessoal,
                identificador=identificador,
                data_admissao=_cell(row, header_index, "Data admissão"),
            )
        )

    return ParsedColaboradoresAdmission(rows=tuple(result), path=str(path))


def parse_solicitacoes_xlsx(path: str | Path) -> ParsedSolicitacoes:
    """Lê ``backup_solicitacoes_*`` e valida colunas obrigatórias do contrato.

    Colunas obrigatórias: ``Identificador``, ``Nome``, ``Iniciada em``,
    ``Terminada em``. ``Status`` é opcional (ignorado na persistência —
    sempre ``encerrado``). ``Criada em`` e demais colunas são ignoradas.
    Linhas sem ``Identificador`` canônico são descartadas. Datas e nome
    permanecem crus para normalização posterior (R6/R7).

    Raises:
        LegacyParseError: arquivo / OOXML / colunas (T025 — fatal).
    """
    path = Path(path)
    raw_rows = _load_sheet_rows(path)
    if not raw_rows:
        raise LegacyParseError(f"Planilha vazia (sem header): {path}")

    header_index = _header_index(raw_rows[0])
    _require_columns(path, header_index, SOLICITACOES_REQUIRED_COLUMNS)

    result: list[SolicitacaoRow] = []
    for excel_row, row in enumerate(raw_rows[1:], start=2):
        identificador = canonicalize_id(_cell(row, header_index, "Identificador"))
        if not identificador:
            continue
        result.append(
            SolicitacaoRow(
                linha=excel_row,
                identificador=identificador,
                nome=_cell(row, header_index, "Nome"),
                iniciada_em=_cell(row, header_index, "Iniciada em"),
                terminada_em=_cell(row, header_index, "Terminada em"),
                status=_as_text(_cell(row, header_index, "Status")),
            )
        )

    return ParsedSolicitacoes(rows=tuple(result), path=str(path))


def parse_avaliacoes_crosswalk_xlsx(path: str | Path) -> ParsedAvaliacoesCrosswalk:
    """Lê apenas colunas de crosswalk de ``backup_avaliacoes_*.xlsx``.

    Demais colunas são ignoradas nesta fatia (column-mapping §avaliações).
    Linhas sem nome ou identificador são descartadas (crosswalk R8).

    Raises:
        LegacyParseError: arquivo / OOXML / colunas (T025 — fatal).
    """
    path = Path(path)
    raw_rows = _load_sheet_rows(path)
    if not raw_rows:
        raise LegacyParseError(f"Planilha vazia (sem header): {path}")

    header_index = _header_index(raw_rows[0])
    _require_columns(path, header_index, AVALIACOES_REQUIRED_COLUMNS)

    result: list[AvaliacaoCrosswalkRow] = []
    for excel_row, row in enumerate(raw_rows[1:], start=2):
        nome = _as_text(_cell(row, header_index, "Nome Avaliado"))
        identificador = canonicalize_id(
            _cell(row, header_index, "Identificador Avaliado")
        )
        if not nome or not identificador:
            continue
        result.append(
            AvaliacaoCrosswalkRow(
                linha=excel_row,
                nome_avaliado=nome,
                identificador_avaliado=identificador,
            )
        )

    return ParsedAvaliacoesCrosswalk(rows=tuple(result), path=str(path))


def parse_avaliacoes_headers_xlsx(path: str | Path) -> ParsedAvaliacoesHeaders:
    """Lê cabeçalhos de ``backup_avaliacoes_*`` para agregação 1:1 (011).

    Colunas obrigatórias: ``Identificador``, ``Identificador Solicitação``,
    ``Identificador Avaliado``, ``Nome Avaliado``, ``Nome Avaliador``.
    Linhas sem qualquer dos três IDs canônicos são descartadas (aggregation
    contract). Demais colunas (ex. ``Avaiação criada em``) são ignoradas.

    Raises:
        LegacyParseError: arquivo / OOXML / colunas (T025 — fatal).
    """
    path = Path(path)
    raw_rows = _load_sheet_rows(path)
    if not raw_rows:
        raise LegacyParseError(f"Planilha vazia (sem header): {path}")

    header_index = _header_index(raw_rows[0])
    _require_columns(path, header_index, AVALIACOES_HEADERS_REQUIRED_COLUMNS)

    result: list[AvaliacaoHeaderRow] = []
    for excel_row, row in enumerate(raw_rows[1:], start=2):
        identificador = canonicalize_id(_cell(row, header_index, "Identificador"))
        solicitacao = canonicalize_id(
            _cell(row, header_index, "Identificador Solicitação")
        )
        avaliado = canonicalize_id(
            _cell(row, header_index, "Identificador Avaliado")
        )
        if not identificador or not solicitacao or not avaliado:
            continue
        result.append(
            AvaliacaoHeaderRow(
                linha=excel_row,
                identificador=identificador,
                identificador_solicitacao=solicitacao,
                identificador_avaliado=avaliado,
                nome_avaliado=_as_text(_cell(row, header_index, "Nome Avaliado")),
                nome_avaliador=_as_text(_cell(row, header_index, "Nome Avaliador")),
            )
        )

    return ParsedAvaliacoesHeaders(rows=tuple(result), path=str(path))


def parse_notas_xlsx(path: str | Path) -> ParsedNotas:
    """Lê ``backup_notas_avaliacoes_*`` (013 — T004).

    Colunas obrigatórias: ``Identificador``, ``Identificador Avaliação``,
    ``Nome Avaliador``, ``Nome Avaliado``, ``Identificador Habilidade``,
    ``habilidade``, ``Fator no Momento``, ``Nota``. Sem coluna de nível.
    Linhas sem ``Identificador Avaliação`` canônico são descartadas.
    ``Fator no Momento`` e ``Nota`` permanecem crus (conflito na persistência).

    Raises:
        LegacyParseError: arquivo / OOXML / colunas (fatal, zero writes).
    """
    path = Path(path)
    raw_rows = _load_sheet_rows(path)
    if not raw_rows:
        raise LegacyParseError(f"Planilha vazia (sem header): {path}")

    header_index = _header_index(raw_rows[0])
    _require_columns(path, header_index, NOTAS_REQUIRED_COLUMNS)

    result: list[NotaRow] = []
    for excel_row, row in enumerate(raw_rows[1:], start=2):
        identificador_avaliacao = canonicalize_id(
            _cell(row, header_index, "Identificador Avaliação")
        )
        if not identificador_avaliacao:
            continue
        result.append(
            NotaRow(
                linha=excel_row,
                identificador=canonicalize_id(
                    _cell(row, header_index, "Identificador")
                ),
                identificador_avaliacao=identificador_avaliacao,
                nome_avaliador=_as_text(_cell(row, header_index, "Nome Avaliador")),
                nome_avaliado=_as_text(_cell(row, header_index, "Nome Avaliado")),
                identificador_habilidade=canonicalize_id(
                    _cell(row, header_index, "Identificador Habilidade")
                ),
                habilidade=_as_text(_cell(row, header_index, "habilidade")),
                fator_no_momento=_cell(row, header_index, "Fator no Momento"),
                nota=_cell(row, header_index, "Nota"),
            )
        )

    return ParsedNotas(rows=tuple(result), path=str(path))


def parse_comentarios_xlsx(path: str | Path) -> ParsedComentarios:
    """Lê ``backup_comentarios_avaliacoes_*`` (013 — T004).

    Colunas obrigatórias: ``Identificador`` (ID de avaliação),
    ``Identificador Avaliador``, ``Nome Avaliador``, ``Nome Avaliado``,
    ``Comentário``, ``Criado em``. Demais colunas do dump (solicitação,
    identificador avaliado) são ignoradas. Linhas sem ``Identificador``
    canônico são descartadas. ``Criado em`` permanece cru para
    ``dates.parse_legacy_datetime``.

    Raises:
        LegacyParseError: arquivo / OOXML / colunas (fatal, zero writes).
    """
    path = Path(path)
    raw_rows = _load_sheet_rows(path)
    if not raw_rows:
        raise LegacyParseError(f"Planilha vazia (sem header): {path}")

    header_index = _header_index(raw_rows[0])
    _require_columns(path, header_index, COMENTARIOS_REQUIRED_COLUMNS)

    result: list[ComentarioRow] = []
    for excel_row, row in enumerate(raw_rows[1:], start=2):
        identificador = canonicalize_id(_cell(row, header_index, "Identificador"))
        if not identificador:
            continue
        result.append(
            ComentarioRow(
                linha=excel_row,
                identificador=identificador,
                identificador_avaliador=canonicalize_id(
                    _cell(row, header_index, "Identificador Avaliador")
                ),
                nome_avaliador=_as_text(_cell(row, header_index, "Nome Avaliador")),
                nome_avaliado=_as_text(_cell(row, header_index, "Nome Avaliado")),
                comentario=_as_text(_cell(row, header_index, "Comentário")),
                criado_em=_cell(row, header_index, "Criado em"),
            )
        )

    return ParsedComentarios(rows=tuple(result), path=str(path))


def parse_habilidades_xlsx(path: str | Path) -> ParsedHabilidades:
    """Lê ``backup_habilidades_*`` opcional (013 — T004; só se ``--habilidades``).

    Colunas: ``Identificador``, ``Habilidade``, ``Grupo``. Linhas sem
    ``Identificador`` canônico são descartadas. **Não** lê a matriz
    ``backup_habilidades_cargo_*``.

    Raises:
        LegacyParseError: arquivo / OOXML / colunas (fatal, zero writes).
    """
    path = Path(path)
    raw_rows = _load_sheet_rows(path)
    if not raw_rows:
        raise LegacyParseError(f"Planilha vazia (sem header): {path}")

    header_index = _header_index(raw_rows[0])
    _require_columns(path, header_index, HABILIDADES_REQUIRED_COLUMNS)

    result: list[HabilidadeRow] = []
    for excel_row, row in enumerate(raw_rows[1:], start=2):
        identificador = canonicalize_id(_cell(row, header_index, "Identificador"))
        if not identificador:
            continue
        result.append(
            HabilidadeRow(
                linha=excel_row,
                identificador=identificador,
                habilidade=_as_text(_cell(row, header_index, "Habilidade")),
                grupo=_as_text(_cell(row, header_index, "Grupo")),
            )
        )

    return ParsedHabilidades(rows=tuple(result), path=str(path))


def _pdi_person_id(row: tuple[Any, ...], header_index: dict[str, int]) -> str:
    """ID Sólides da pessoa se o header existir (não fatal se ausente).

    Preferência: ``Identificador``; senão ``Identificador Avaliado``.
    """
    if "Identificador" in header_index:
        return canonicalize_id(_cell(row, header_index, "Identificador"))
    if "Identificador Avaliado" in header_index:
        return canonicalize_id(_cell(row, header_index, "Identificador Avaliado"))
    return ""


def parse_pdi_xlsx(path: str | Path) -> ParsedPdi:
    """Lê ``backup_pdi_*`` (014 — T004).

    Colunas obrigatórias: ``Nome``, ``Título do PDI``, ``Status``,
    ``Objetivo``, ``Situação Atual``, ``Situação Desejada``,
    ``Data de Entrega``. Opcionais se o header existir: ``Criado em``,
    ``Identificador Solicitação``, ID da pessoa (``Identificador`` /
    ``Identificador Avaliado``). Células de data/texto permanecem cruas.
    Células vazias **não** são erro de parse.

    Raises:
        LegacyParseError: arquivo / OOXML / colunas (fatal, zero writes).
    """
    path = Path(path)
    raw_rows = _load_sheet_rows(path)
    if not raw_rows:
        raise LegacyParseError(f"Planilha vazia (sem header): {path}")

    header_index = _header_index(raw_rows[0])
    _require_columns(path, header_index, PDI_REQUIRED_COLUMNS)

    result: list[PdiRow] = []
    for excel_row, row in enumerate(raw_rows[1:], start=2):
        result.append(
            PdiRow(
                linha=excel_row,
                nome=_cell(row, header_index, "Nome"),
                titulo=_cell(row, header_index, "Título do PDI"),
                status=_cell(row, header_index, "Status"),
                objetivo=_cell(row, header_index, "Objetivo"),
                situacao_atual=_cell(row, header_index, "Situação Atual"),
                situacao_desejada=_cell(row, header_index, "Situação Desejada"),
                data_entrega=_cell(row, header_index, "Data de Entrega"),
                criado_em=_cell(row, header_index, "Criado em"),
                identificador_solicitacao=canonicalize_id(
                    _cell(row, header_index, "Identificador Solicitação")
                ),
                identificador_pessoa=_pdi_person_id(row, header_index),
            )
        )

    return ParsedPdi(rows=tuple(result), path=str(path))
